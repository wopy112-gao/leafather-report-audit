#!/usr/bin/env python3
"""rla_extract.py — 从报告里抽出所有数字，并判断它是不是"指标数字"。

支持：md / txt / html / docx / pptx / xlsx
输出：JSON 列表，每条带 值、写法、上下文、分类（metric / date / ordinal / year / misc）

为什么要有"分类"这一层：
    报告里除了指标，还有日期、页码、条款号、"3 条建议"。
    不分就淹在误报里，没人会用。分类靠上下文词 + 数字形态两路判断。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

NUM_RE = re.compile(r"\d[\d,，]*(?:\.\d+)?\s*(?:%|％|万|亿|千)?")
DATE_PATTERNS = [
    re.compile(r"\d{4}\s*[-/年]\s*\d{1,2}\s*[-/月]\s*\d{1,2}\s*日?"),
    re.compile(r"\d{4}\s*[-/年]\s*\d{1,2}\s*月?"),
    re.compile(r"\d{1,2}\s*月\s*\d{1,2}\s*日"),
]
ORDINAL_RE = re.compile(r"(第\s*\d+\s*[章节条款页图表篇]|^\s*\d+[\.、)]|表\s*\d+|图\s*\d+|P\d+|§\d+)")
METRIC_HINT = re.compile(
    r"(场景|门店|店员|客户|单|额|元|万|盒|个|家|人|率|占比|份额|增速|数量|总数|覆盖|渗透|客单|环比|同比"
    # 2026-10-25 补：这些都是"数据量词"，报告里极常见，漏掉会被判 misc 静默豁免
    # （实测：一份用「总行数 / 成交行数」写法的报告，3 个真实指标全被豁免，结论还显示通过）
    r"|行数|记录数|条数|笔数|份数|件数|次数|明细数"
    r")"
)
YEAR_RE = re.compile(r"^(19|20)\d{2}$")


def is_identifier(text: str, start: int, end: int) -> bool:
    """这个数字是不是标识符/编号的一部分（不是数据）。

    实测：某 agent那份报告里 306 个「指标」有 40 个是 `GLP-1` 里的 `1`。
    这类误抽会同时污染「可追溯率」和「无法判定数」两个指标，必须先清掉。

    认这几种：
      - 紧跟英文字母：GLP-1 / L1 / V2 / S31
      - 紧跟「字母-」：GLP-1
      - 序号：Step 5 / 第 5
      - 章节号：行首或井号后的 5.1
    """
    left = text[max(0, start - 6):start]
    # NUM_RE 会把尾随空格吞进 match（"5.1 "），必须 strip 才能判断章节号
    token = text[start:end].strip()
    if re.search(r"[A-Za-z][-]?$", left):          # GLP-1 / L1 / S31
        return True
    if re.search(r"(Step|step|第)\s*$", left):      # Step 5 / 第 5
        return True
    if re.fullmatch(r"\d+\.\d+", token) and re.match(r"^\s*#{0,6}\s*$", left):
        return True                                  # 5.1 章节号
    return False


def norm(raw: str) -> tuple[float | None, str]:
    """把 '15,597' / '1.56万' / '73.9%' 归一成 (数值, 单位提示)。"""
    s = raw.strip().replace("，", ",").replace(" ", "")
    unit = ""
    if s.endswith(("%", "％")):
        unit, s = "pct1", s[:-1]
    for u, mul in (("亿", 1e8), ("万", 1e4), ("千", 1e3)):
        if s.endswith(u):
            unit, s = u, s[:-1]
            break
    s = s.replace(",", "")
    try:
        v = float(s)
    except ValueError:
        return None, unit
    if unit == "亿":
        v *= 1e8
    elif unit == "万":
        v *= 1e4
    elif unit == "千":
        v *= 1e3
    return v, unit


def looks_like_data(raw: str) -> bool:
    """看着像不像一个数据（而不是行文里的数字）。

    2026-10-26（方案 B）：默认立场倒过来 —— 白名单永远不全（补了「行数」，人家写「样本池」），
    所以**带千分位 / 百分比 / 上千的数字，先当数据、要求它给出处**，让举证责任回到报告方。
    实测依据：同一份报告只把「总行数」改成「样本池规模」，4 个指标全被豁免、结论照样显示通过。
    """
    s = (raw or "").strip()
    if not s:
        return False
    if s.endswith(("%", "％")):
        return True
    if "," in s or "，" in s:                 # 带千分位 → 基本不可能是行文小数字
        return True
    core = s.rstrip("%％万亿千").replace(",", "").replace("，", "")
    try:
        return abs(float(core)) >= 1000
    except ValueError:
        return False


def classify(raw: str, ctx: str) -> str:
    if any(p.search(ctx) for p in DATE_PATTERNS):
        return "date"
    if ORDINAL_RE.search(ctx):
        return "ordinal"
    core = raw.strip().rstrip("%％万亿千").replace(",", "").replace("，", "")
    if YEAR_RE.match(core) and not METRIC_HINT.search(ctx):
        return "year"
    if METRIC_HINT.search(ctx):
        return "metric"
    # 方案 B：词表没命中，但数字本身长得就像数据 → 仍然要追（举证责任在报告方）
    if looks_like_data(raw):
        return "metric"
    return "misc"


SKIP_START = "<!-- audit-skip:start -->"
SKIP_END = "<!-- audit-skip:end -->"
SKIP_NUMS = {"n": 0}          # 被跳过的数字个数（供结论透明披露）


def strip_skip(text: str) -> str:
    """剔除 <!-- audit-skip --> 区（纯技术说明）。

    透明原则：跳过多少数字必须统计出来写进结论，不能偷偷绕过。
    """
    out, i = [], 0
    while True:
        a = text.find(SKIP_START, i)
        if a < 0:
            out.append(text[i:])
            break
        out.append(text[i:a])
        b = text.find(SKIP_END, a)
        if b < 0:
            out.append(text[a:])
            break
        skipped = text[a:b]
        SKIP_NUMS["n"] += len(NUM_RE.findall(skipped))
        i = b + len(SKIP_END)
    return "".join(out)


def iter_text(path: Path, warnings: list | None = None):
    """逐段吐 (文本, 位置标签)。图片等扫不到的内容写进 warnings。"""
    suf = path.suffix.lower()
    if suf in (".md", ".txt", ".markdown", ""):
        raw = strip_skip(path.read_text(encoding="utf-8", errors="ignore"))
        for i, line in enumerate(raw.splitlines(), 1):
            yield line, f"L{i}"
    elif suf in (".html", ".htm"):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(strip_skip(path.read_text(encoding="utf-8", errors="ignore")), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        for i, line in enumerate(soup.get_text("\n").splitlines(), 1):
            if line.strip():
                yield line.strip(), f"L{i}"
    elif suf == ".docx":
        try:
            import docx
        except ImportError:
            raise SystemExit("❌ 缺 python-docx（pip install python-docx）")
        d = docx.Document(str(path))
        for i, p in enumerate(d.paragraphs, 1):
            if p.text.strip():
                yield p.text, f"P{i}"
        for ti, t in enumerate(d.tables, 1):
            for ri, row in enumerate(t.rows, 1):
                cells = [c.text.strip() for c in row.cells]
                yield " | ".join(cells), f"T{ti}R{ri}"
    elif suf == ".pptx":
        try:
            from pptx import Presentation
        except ImportError:
            raise SystemExit("❌ 缺 python-pptx（pip install python-pptx）")
        prs = Presentation(str(path))
        pics = 0

        def _is_picture(sh) -> bool:
            """python-pptx 的 shape_type 是枚举，不能用 == 13 直接比（踩过）。"""
            try:
                from pptx.enum.shapes import MSO_SHAPE_TYPE
                return sh.shape_type == MSO_SHAPE_TYPE.PICTURE
            except Exception:
                return "PICTURE" in str(getattr(sh, "shape_type", ""))

        def walk(shapes, tag):
            nonlocal pics
            for sh in shapes:
                if _is_picture(sh):
                    pics += 1
                if sh.has_text_frame and sh.text_frame.text.strip():
                    yield sh.text_frame.text, tag
                if getattr(sh, "has_table", False) and sh.has_table:
                    for ri, row in enumerate(sh.table.rows, 1):
                        yield " | ".join(c.text.strip() for c in row.cells), f"{tag}-T{ri}"
                if "GROUP" in str(getattr(sh, "shape_type", "")) or getattr(sh, "shape_type", None) == 6:
                    yield from walk(sh.shapes, tag)

        for si, slide in enumerate(prs.slides, 1):
            yield from walk(slide.shapes, f"S{si}")
        if pics:
            msg = f"{path.name}：含 {pics} 张图片，图内数字无法扫描（需人工核）"
            if warnings is not None:
                warnings.append(msg)
            else:
                print(f"⚠️  {msg}", file=sys.stderr)
    elif suf == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            raise SystemExit("❌ 缺 openpyxl（pip install openpyxl）")
        wb = openpyxl.load_workbook(str(path), data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    yield str(c.value), f"{ws.title}!{c.coordinate}"
    else:
        raise SystemExit(f"❌ 不支持的格式：{suf}")


def extract_with_warnings(path: Path) -> tuple[list[dict], list[str]]:
    """抽取数字，并返回"扫不到的内容"提示（图片等）。"""
    warnings: list[str] = []
    out = []
    for text, loc in iter_text(path, warnings):
        for m in NUM_RE.finditer(text):
            raw = m.group(0).strip()
            if not raw or not any(ch.isdigit() for ch in raw):
                continue
            if is_identifier(text, m.start(), m.end()):
                continue
            v, unit = norm(raw)
            if v is None:
                continue
            left = max(0, m.start() - 24)
            ctx = text[left:m.end() + 24]
            out.append({
                "file": path.name,
                "loc": loc,
                "raw": raw,
                "value": v,
                "unit_hint": unit,
                "kind": classify(raw, ctx),
                "context": ctx.strip(),
            })
    if SKIP_NUMS["n"]:
        warnings.append(f"{path.name}：技术说明区（audit-skip）跳过 {SKIP_NUMS['n']} 个数字，未审计")
        SKIP_NUMS["n"] = 0
    return out, warnings


def extract(path: Path) -> list[dict]:
    """兼容入口：只要数字，丢弃提示。"""
    numbers, _ = extract_with_warnings(path)
    return numbers


def main() -> int:
    ap = argparse.ArgumentParser(description="从报告里抽数字")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--json", action="store_true", help="输出 JSON")
    ap.add_argument("--only-metric", action="store_true", help="只列指标类")
    args = ap.parse_args()

    all_numbers = []
    for f in args.files:
        p = Path(f)
        if not p.exists():
            print(f"❌ 文件不存在：{f}", file=sys.stderr)
            return 2
        all_numbers.extend(extract(p))

    if args.only_metric:
        all_numbers = [n for n in all_numbers if n["kind"] == "metric"]

    if args.json:
        print(json.dumps(all_numbers, ensure_ascii=False, indent=2))
        return 0

    from collections import Counter
    c = Counter(n["kind"] for n in all_numbers)
    print(f"共抽到 {len(all_numbers)} 个数字｜" + "  ".join(f"{k}={v}" for k, v in c.most_common()))
    for n in all_numbers:
        if n["kind"] == "metric":
            print(f"  [{n['file']} {n['loc']}] {n['raw']:>12}  ← {n['context'][:56]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
